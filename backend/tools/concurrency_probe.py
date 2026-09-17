"""Disposable-environment probes. Never import/register these tasks in production."""
# ruff: noqa: E402

import asyncio
import io
import json
import os
import time
import uuid

import httpx
from redis import Redis
from sqlalchemy import select, update

if os.environ.get("KAP_ACCEPTANCE_ISOLATED") != "1":
    raise RuntimeError("Only run with deploy/compose.concurrency-test.yml")
if (
    os.environ.get("DATABASE_URL")
    != "postgresql+asyncpg://test:test@postgres:5432/kap_concurrency_test"
):
    raise RuntimeError("Refusing unexpected database endpoint")
if os.environ.get("REDIS_URL") != "redis://redis:6379/0":
    raise RuntimeError("Refusing unexpected broker endpoint")

from app.core.config import get_settings
from app.db.base import Base
from app.db.session import get_engine
from app.models.ingest import IngestTask
from app.models.knowledge import KnowledgeAsset, KnowledgeAssetVersion
from app.models.naming import NamingRuleRevision
from app.seed.dev_seed import USER_BOSS, USER_CONSULTANT, USER_PROJECT_MANAGER, seed_dev_identities
from app.services.directories import default_directory_config
from app.services.ingest_capacity import processing_slot
from app.services.jobs.parse_reconcile import _claim
from app.worker.celery_app import celery_app
from app.worker.queues import ingest_processing_queue
from app.worker.runtime import run_task


def broker():
    return Redis.from_url(get_settings().redis_url)


@celery_app.task(name="acceptance.hold")
def hold(run, seconds):
    broker().incr(f"{run}:started")
    time.sleep(seconds)
    return os.getpid()


@celery_app.task(name="acceptance.ping")
def ping():
    return os.getpid()


@celery_app.task(name="acceptance.slot")
def slot(task_id):
    async def work(maker):
        async with processing_slot(maker, uuid.UUID(task_id)) as admitted:
            if admitted:
                await asyncio.sleep(3)
            return admitted

    return run_task(work)


@celery_app.task(name="acceptance.claim")
def claim(version_id, run, delay=0):
    async def work(maker):
        async with maker() as session:
            snapshot = (
                await session.execute(
                    select(
                        KnowledgeAssetVersion.id,
                        KnowledgeAssetVersion.parse_reconcile_until,
                    ).where(KnowledgeAssetVersion.id == uuid.UUID(version_id))
                )
            ).one()
            await session.commit()
            broker().incr(f"{run}:claim-ready")
            deadline = time.monotonic() + 15
            while int(broker().get(f"{run}:claim-ready") or 0) < 4:
                if time.monotonic() > deadline:
                    raise TimeoutError("claim barrier")
                await asyncio.sleep(0.05)
            # One contender uses the same snapshot after the winner released it.
            await asyncio.sleep(delay)
            async with _claim(session, snapshot) as token:
                if token:
                    await asyncio.sleep(1)
                return token is not None

    return run_task(work)


async def seed():
    engine = get_engine()
    if engine.url.database != "kap_concurrency_test":
        raise RuntimeError("Refusing non-acceptance database")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    from sqlalchemy.ext.asyncio import async_sessionmaker

    maker = async_sessionmaker(engine, expire_on_commit=False)
    async with maker() as session:
        await seed_dev_identities(session)
        if await session.scalar(select(NamingRuleRevision.id).limit(1)) is None:
            session.add(
                NamingRuleRevision(
                    version=1,
                    base_published_version=0,
                    status="published",
                    config={
                        "schema_version": 2,
                        "enforced": True,
                        "project_codes": [],
                        "categories": [],
                        "directories": default_directory_config(),
                    },
                )
            )
        tasks = [
            IngestTask(
                source="path_b_upload",
                source_file_ref="acceptance/unused",
                source_file_name=f"sample-{i}.txt",
                source_file_hash=f"{i:064x}",
                created_by=USER_BOSS,
                target_scope="company",
                status="pending_confirmation",
            )
            for i in range(100)
        ]
        heavy = IngestTask(
            source="path_b_upload",
            source_file_ref="acceptance/unused",
            source_file_name="sample.pdf",
            created_by=USER_BOSS,
            status="processing",
        )
        asset = KnowledgeAsset(
            title="Acceptance",
            scope="company",
            asset_type="document",
            owner_user_id=USER_BOSS,
            confidentiality_level="L2",
        )
        session.add_all([*tasks, heavy, asset])
        await session.flush()
        version = KnowledgeAssetVersion(
            asset_id=asset.id, version_no="V1", version_status="active", created_by=USER_BOSS
        )
        session.add(version)
        await session.commit()
        result = [str(task.id) for task in tasks], str(heavy.id), str(version.id)
    await engine.dispose()
    return result


async def preview_load(task_ids):
    latencies = []
    gate = asyncio.Semaphore(4)
    async with httpx.AsyncClient(base_url="http://api:8000", timeout=30) as client:

        async def request(index):
            async with gate:
                start = time.monotonic()
                ids = task_ids[(index % 4) * 25 : (index % 4 + 1) * 25]
                response = await client.post(
                    "/api/v1/ingest/bulk-naming-preview",
                    headers={"X-Dev-User-Id": str(USER_BOSS)},
                    json={
                        "target_scope": "company",
                        "items": [
                            {
                                "task_id": task_id,
                                "confidentiality_level": "L2",
                                "naming": {
                                    "directory_key": "company.methodology",
                                    "subject": "Acceptance",
                                    "formed_on": "2026-09-17",
                                    "version": "V1",
                                    "applicable_to": "通用",
                                },
                            }
                            for task_id in ids
                        ],
                    },
                )
                response.raise_for_status()
                assert all(item["submittable"] for item in response.json()["items"])
                latencies.append(time.monotonic() - start)

        await asyncio.gather(*(request(i) for i in range(40)))
    latencies.sort()
    return {
        "requests": len(latencies),
        "concurrency": 4,
        "p95_seconds": latencies[37],
        "max_seconds": latencies[-1],
    }


async def finish_capacity_probe(task_id):
    # The synthetic running task must not consume scheduling capacity for uploads.
    engine = get_engine()
    async with engine.begin() as connection:
        await connection.execute(
            update(IngestTask)
            .where(IngestTask.id == uuid.UUID(task_id))
            .values(status="cancelled", cancel_requested=True)
        )
    await engine.dispose()


def document_bytes(kind, label):
    buffer = io.BytesIO()
    text = (
        f"Synthetic consulting acceptance document {label}. "
        "This document contains test content, not customer data. " * 30
    )
    if kind == "txt":
        return text.encode(), "text/plain"
    if kind == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph(text)
        document.save(buffer)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif kind == "pptx":
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = label
        slide.placeholders[1].text = text
        presentation.save(buffer)
        mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif kind == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append([label, text])
        workbook.save(buffer)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        import pymupdf

        document = pymupdf.open()
        page = document.new_page()
        page.insert_textbox(pymupdf.Rect(40, 40, 550, 800), text, fontsize=10)
        return document.tobytes(), "application/pdf"
    return buffer.getvalue(), mime


async def upload_load():
    gate = asyncio.Semaphore(5)
    elapsed = []
    async with httpx.AsyncClient(base_url="http://api:8000", timeout=60) as client:

        async def upload(user_id, kind):
            async with gate:
                start = time.monotonic()
                label = f"acceptance-{uuid.uuid4()}"
                content, mime = document_bytes(kind, label)
                headers = {"X-Dev-User-Id": str(user_id)}
                response = await client.post(
                    "/api/v1/ingest/upload",
                    headers=headers,
                    data={"target_scope": "personal", "formed_on": "2026-09-17"},
                    files={"file": (f"{label}.{kind}", content, mime)},
                )
                response.raise_for_status()
                task_id = response.json()["ingest_task_id"]
                deadline = time.monotonic() + 240
                while True:
                    status_response = await client.get(
                        f"/api/v1/ingest/{task_id}/status", headers=headers
                    )
                    status_response.raise_for_status()
                    status = status_response.json()
                    if status["stage"] == "awaiting_confirmation":
                        break
                    if status["status"] == "failed" or time.monotonic() > deadline:
                        raise AssertionError(
                            f"{kind}: {status.get('stage')} / {status.get('error')}"
                        )
                    await asyncio.sleep(0.5)
                confirmed = await client.post(
                    f"/api/v1/ingest/{task_id}/confirm",
                    headers=headers,
                    json={
                        "title": label,
                        "summary": "Synthetic acceptance material",
                        "tags": [],
                        "target_scope": "personal",
                        "target_zone": "material",
                        "confidentiality_level": "L2",
                        "directory_key": "personal.learning_notes",
                    },
                )
                confirmed.raise_for_status()
                assert confirmed.json()["status"] == "completed"
                assert confirmed.json()["result_asset_id"]
                elapsed.append(time.monotonic() - start)

        await asyncio.gather(
            *(
                upload(user, kind)
                for user in [USER_BOSS, USER_CONSULTANT, USER_PROJECT_MANAGER]
                for kind in ["txt", "docx", "pptx", "xlsx", "pdf"]
            )
        )
    elapsed.sort()
    return {
        "users": 3,
        "files": len(elapsed),
        "concurrency": 5,
        "p95_seconds": elapsed[-1],
        "llm": "deterministic test stand-in",
        "weknora": "disabled",
    }


async def transport_load():
    confirmation_gate = asyncio.Semaphore(5)
    async with httpx.AsyncClient(base_url="http://api:8000", timeout=60) as client:

        async def run_session(user_id):
            start = time.monotonic()
            session_id = str(uuid.uuid4())
            headers = {"X-Dev-User-Id": str(user_id)}
            contents = [f"Synthetic batch {session_id} item {i}. " * 30 for i in range(10)]
            manifest = [
                {
                    "client_file_key": str(i),
                    "file_name": f"{session_id}-{i}.txt",
                    "file_size": len(content.encode()),
                    "file_type": "text/plain",
                    "formed_on": "2026-09-17",
                    "transport_batch_index": i // 5,
                }
                for i, content in enumerate(contents)
            ]
            initialized = await client.post(
                "/api/v1/ingest/upload-sessions/init",
                headers=headers,
                json={
                    "session_id": session_id,
                    "target_scope": "personal",
                    "total_transport_batches": 2,
                    "manifest": manifest,
                },
            )
            initialized.raise_for_status()
            item_ids = [item["id"] for item in initialized.json()["items"]]
            base = f"/api/v1/ingest/upload-sessions/{session_id}"
            for batch in range(2):
                indices = list(range(batch * 5, batch * 5 + 5))
                data = {
                    "batch_id": f"batch-{batch}",
                    "batch_index": str(batch),
                    "item_ids": json.dumps([item_ids[i] for i in indices]),
                }
                files = [
                    ("files", (manifest[i]["file_name"], contents[i].encode(), "text/plain"))
                    for i in indices
                ]
                # Identical requests in flight at once must not duplicate tasks or bytes.
                replies = await asyncio.gather(
                    *(
                        client.post(f"{base}/batches", headers=headers, data=data, files=files)
                        for _ in range(2)
                    )
                )
                for reply in replies:
                    reply.raise_for_status()
                    assert reply.json()["uploaded_files"] == (batch + 1) * 5
                assert [item["ingest_task_id"] for item in replies[0].json()["items"]] == [
                    item["ingest_task_id"] for item in replies[1].json()["items"]
                ]
            completed = await client.post(f"{base}/complete", headers=headers)
            completed.raise_for_status()
            assert completed.json()["upload_completed"]
            deadline = time.monotonic() + 240
            while True:
                response = await client.get(base, headers=headers)
                response.raise_for_status()
                state = response.json()
                assert state["total_files"] == len(state["items"]) == 10
                assert not state["failed_files"], state
                if all(item["status"] == "awaiting_confirmation" for item in state["items"]):
                    break
                if time.monotonic() > deadline:
                    raise TimeoutError("batch processing did not complete")
                await asyncio.sleep(0.5)

            async def confirm(item):
                async with confirmation_gate:
                    response = await client.post(
                        f"/api/v1/ingest/{item['ingest_task_id']}/confirm",
                        headers=headers,
                        json={
                            "title": item["file_name"],
                            "summary": "Synthetic",
                            "target_scope": "personal",
                            "confidentiality_level": "L2",
                            "directory_key": "personal.learning_notes",
                        },
                    )
                    response.raise_for_status()
                    assert response.json()["status"] == "completed"
                    return response.json()["result_asset_id"]

            asset_ids = await asyncio.gather(*(confirm(item) for item in state["items"]))
            assert len(set(asset_ids)) == 10 and None not in asset_ids
            return time.monotonic() - start

        times = await asyncio.gather(
            *(run_session(user) for user in [USER_BOSS, USER_CONSULTANT, USER_PROJECT_MANAGER])
        )
    return {
        "sessions": 3,
        "files": 30,
        "concurrent_same_batch_replays": 6,
        "max_session_seconds": max(times),
        "confirmed": 30,
    }


async def http_load(task_ids):
    preview, upload, transport = await asyncio.gather(
        preview_load(task_ids), upload_load(), transport_load()
    )
    return {"preview_load": preview, "upload_load": upload, "transport_load": transport}


def main():
    task_ids, heavy_id, version_id = asyncio.run(seed())
    run = f"acceptance:{uuid.uuid4()}"
    settings = get_settings()
    queue = ingest_processing_queue(settings, file_name="sample.pdf", mime_type="application/pdf")
    holds = [hold.apply_async((run, 8), queue=queue) for _ in range(4)]
    deadline = time.monotonic() + 30
    while int(broker().get(f"{run}:started") or 0) < 4:
        if time.monotonic() > deadline:
            raise TimeoutError("workers did not start")
        time.sleep(0.1)
    start = time.monotonic()
    ping.apply_async(queue=settings.celery_default_queue).get(timeout=5)
    light_latency = time.monotonic() - start
    assert not all(task.ready() for task in holds), "light task waited for heavy queue"
    assert len({task.get(timeout=20) for task in holds}) == 4
    slots = [slot.apply_async((heavy_id,), queue=queue) for _ in range(4)]
    assert sum(task.get(timeout=20) for task in slots) == 1
    assert slot.apply_async((heavy_id,), queue=queue).get(timeout=20) is True
    claims = [claim.apply_async((version_id, run, delay), queue=queue) for delay in [0, 0, 0, 2]]
    assert sum(task.get(timeout=30) for task in claims) == 1
    asyncio.run(finish_capacity_probe(heavy_id))
    report = {
        "heavy_queue_isolation": "passed",
        "postgres_capacity": "passed",
        "postgres_claim": "passed",
        "light_latency_seconds": light_latency,
        **asyncio.run(http_load(task_ids)),
    }
    print(json.dumps(report))


if __name__ == "__main__":
    main()
