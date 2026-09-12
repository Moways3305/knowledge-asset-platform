"""Legacy conversion isolation, failure taxonomy, and optional real-engine smoke."""

from __future__ import annotations

import io
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.services import legacy_office
from app.services.extraction import _ControlledExtractionError, _extract_unbounded, extract_text

OLE_HEADER = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@pytest.mark.parametrize("extension", ["doc", "ppt"])
def test_conversion_uses_private_profile_and_cleans_files(monkeypatch, extension):
    monkeypatch.setattr(legacy_office.shutil, "which", lambda _: "libreoffice")
    roots = []

    def run(args, **kwargs):
        root = Path(args[-1]).parent
        roots.append(root)
        assert (root / f"source.{extension}").read_bytes() == OLE_HEADER
        assert "MacroSecurityLevel" in (root / "profile/user/registrymodifications.xcu").read_text()
        assert f"-env:UserInstallation={(root / 'profile').as_uri()}" in args
        assert kwargs["timeout"] == 30
        assert kwargs["stdin"] == subprocess.DEVNULL
        assert kwargs["stdout"] == subprocess.DEVNULL
        (root / f"source.{extension}x").write_bytes(b"converted")
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(legacy_office.subprocess, "run", run)
    assert legacy_office.convert_legacy_office(OLE_HEADER, extension) == (
        b"converted",
        extension + "x",
    )
    assert all(not root.exists() for root in roots)


@pytest.mark.parametrize(
    "failure,code",
    [
        ("timeout", "office_conversion_timeout"),
        ("missing", "office_conversion_failed"),
        ("large", "extraction_structure_limit"),
    ],
)
def test_conversion_failures_are_safe_and_cleanup(monkeypatch, failure, code):
    monkeypatch.setattr(legacy_office.shutil, "which", lambda _: "libreoffice")
    monkeypatch.setattr(legacy_office, "MAX_CONVERTED_BYTES", 2)
    roots = []

    def run(args, **kwargs):
        roots.append(Path(args[-1]).parent)
        if failure == "timeout":
            raise subprocess.TimeoutExpired("private source", 30)
        if failure == "large":
            (roots[-1] / "source.docx").write_bytes(b"123")
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(legacy_office.subprocess, "run", run)
    with pytest.raises(_ControlledExtractionError) as error:
        legacy_office.convert_legacy_office(OLE_HEADER, "doc")
    assert error.value.code == code
    assert "private" not in error.value.message
    assert all(not root.exists() for root in roots)


def test_rtf_doc_reuses_docx_parser(monkeypatch):
    from docx import Document

    document = Document()
    document.add_paragraph("Legacy office body")
    buffer = io.BytesIO()
    document.save(buffer)
    monkeypatch.setattr(
        legacy_office, "convert_legacy_office", lambda *_: (buffer.getvalue(), "docx")
    )
    result = _extract_unbounded(b"{\\rtf1 body}", file_name="original.doc", mime=None)
    assert result.status == "extracted"
    assert "Legacy office body" in result.text


def test_missing_converter_is_not_file_corruption(monkeypatch):
    monkeypatch.setattr(legacy_office.shutil, "which", lambda _: None)
    monkeypatch.setattr(legacy_office.Path, "is_file", lambda _: False)
    with pytest.raises(_ControlledExtractionError) as error:
        legacy_office.convert_legacy_office(OLE_HEADER, "ppt")
    assert error.value.code == "office_converter_unavailable"


@pytest.mark.parametrize(
    "extension,mime", [("doc", "application/msword"), ("ppt", "application/vnd.ms-powerpoint")]
)
def test_legacy_recovery_uses_heavy_queue(extension, mime):
    from app.core.config import get_settings
    from app.models.ingest import IngestTask
    from app.services.jobs.ingest_recovery import _recovery_queue

    task = IngestTask(source_file_name=f"source.{extension}", source_file_mime_type=mime)
    assert _recovery_queue(task) == get_settings().celery_ocr_queue
    task.processing_stage = "content_generation"
    assert _recovery_queue(task) == get_settings().celery_default_queue


@pytest.mark.parametrize("extension", ["doc", "ppt"])
async def test_legacy_first_dispatch_uses_heavy_queue(client, db_session, monkeypatch, extension):
    from app.core.config import get_settings
    from app.models.ingest import IngestTask
    from app.seed.dev_seed import USER_CONSULTANT
    from app.worker import enqueue
    from app.worker.tasks.ingest import process_ingest_upload

    settings = get_settings().model_copy(update={"celery_task_always_eager": False})
    monkeypatch.setattr(enqueue, "get_settings", lambda: settings)
    calls = []
    monkeypatch.setattr(process_ingest_upload, "apply_async", lambda **kwargs: calls.append(kwargs))
    task = IngestTask(
        source="path_b_upload",
        source_file_ref="test-only",
        source_file_name=f"sample.{extension}",
        status="processing",
        created_by=USER_CONSULTANT,
    )
    db_session.add(task)
    await db_session.commit()
    await enqueue.enqueue_ingest_processing(
        db_session,
        task.id,
        storage=client._kap_storage,
        llm=None,
        desensitizer=None,
        trace_id=None,
    )
    assert calls[0]["queue"] == settings.celery_ocr_queue


@pytest.mark.parametrize("extension", ["doc", "ppt"])
def test_real_legacy_round_trip(tmp_path, extension):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        candidate = Path("C:/Program Files/LibreOffice/program/soffice.com")
        executable = str(candidate) if candidate.is_file() else None
    if executable is None:
        pytest.skip("LibreOffice is not installed; install production image dependencies for smoke")
    modern = tmp_path / f"sample.{extension}x"
    if extension == "doc":
        from docx import Document

        document = Document()
        document.add_paragraph("Legacy conversion verification")
        document.save(str(modern))
    else:
        from pptx import Presentation

        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[0])
        slide.shapes.title.text = "Legacy conversion verification"
        document.save(str(modern))
    result = subprocess.run(
        [
            executable,
            f"-env:UserInstallation={(tmp_path / 'synthetic-profile').as_uri()}",
            "--headless",
            "--convert-to",
            "doc:MS Word 97" if extension == "doc" else "ppt:MS PowerPoint 97",
            "--outdir",
            str(tmp_path),
            str(modern),
        ],
        capture_output=True,
        timeout=40,
    )
    assert result.returncode == 0
    original = (tmp_path / f"sample.{extension}").read_bytes()
    parsed = extract_text(original, file_name=f"sample.{extension}", mime=None)
    assert parsed.status == "extracted", parsed
    assert "Legacy conversion verification" in parsed.text
    assert (tmp_path / f"sample.{extension}").read_bytes() == original
